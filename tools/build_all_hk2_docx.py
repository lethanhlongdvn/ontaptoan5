# -*- coding: utf-8 -*-
"""
Script xuất toàn bộ 816 câu hỏi Học kỳ 2 (Trạm 19 - 35) và Trọn bộ 35 Trạm ra file Word (.docx)
Định dạng bảng chuẩn 3 cột: TT | Đề | Đáp án đúng
Xuất:
1. File tổng hợp HK2: Bo_De_Toan_5_Tuan_19_den_35.docx
2. 17 file riêng lẻ HK2: De_Tuan_19_Tram_19.docx -> De_Tuan_35_Tram_35.docx
3. File Đại tổng hợp 35 Trạm (1 680 câu): Bo_De_Toan_5_Tron_Bo_35_Tuan.docx
Lưu tại thư mục: c:\\Users\\Admin\\Desktop\\EduBot-BTCT5\\Đề
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from data_tuan_01_06 import TUAN_01_TO_06
from data_tuan_07_12 import TUAN_07_TO_12
from data_tuan_13_18 import TUAN_13_TO_18
from data_tuan_19_24 import TUAN_19_TO_24
from data_tuan_25_30 import TUAN_25_TO_30
from data_tuan_31_35 import TUAN_31_TO_35

STATIONS_HK1 = TUAN_01_TO_06 + TUAN_07_TO_12 + TUAN_13_TO_18
STATIONS_HK2 = TUAN_19_TO_24 + TUAN_25_TO_30 + TUAN_31_TO_35
ALL_35_STATIONS = STATIONS_HK1 + STATIONS_HK2

def setup_document_styles(subtitle="HỌC KỲ 2: TUẦN 19 - 35"):
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.0)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"EDUBOT - BỘ ĐỀ TOÁN 5 XUYÊN VIỆT ({subtitle})")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(120, 120, 120)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("Trang ")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(9)
        frun.font.color.rgb = RGBColor(100, 100, 100)

        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        fp._p.append(fldSimple)

        frun2 = fp.add_run(" / ")
        frun2.font.name = "Times New Roman"
        frun2.font.size = Pt(9)
        frun2.font.color.rgb = RGBColor(100, 100, 100)

        fldSimple_n = OxmlElement('w:fldSimple')
        fldSimple_n.set(qn('w:instr'), 'NUMPAGES')
        fp._p.append(fldSimple_n)

    return doc

def set_cell_border(cell, color="CBD5E0", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="single" w:sz="{}" w:space="0" w:color="{}"/>
            <w:left w:val="single" w:sz="{}" w:space="0" w:color="{}"/>
            <w:bottom w:val="single" w:sz="{}" w:space="0" w:color="{}"/>
            <w:right w:val="single" w:sz="{}" w:space="0" w:color="{}"/>
        </w:tcBorders>
    '''.format(nsdecls('w'), sz, color, sz, color, sz, color, sz, color))
    tcPr.append(borders)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color_hex))
    tcPr.append(shd)

def add_section_header_row(table, title, fill_color, text_color=(26, 54, 93)):
    row = table.add_row()
    a = row.cells[0]
    b = row.cells[1]
    c = row.cells[2]
    a.merge(b).merge(c)
    
    p = a.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(*text_color)
    
    set_cell_shading(a, fill_color)
    set_cell_border(a, color="A0AEC0", sz="6")

def add_sub_group_header_row(table, sub_title, fill_color="EDF2F7", text_color=(74, 85, 104)):
    row = table.add_row()
    a = row.cells[0]
    b = row.cells[1]
    c = row.cells[2]
    a.merge(b).merge(c)

    p = a.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(sub_title)
    run.font.name = "Times New Roman"
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(*text_color)

    set_cell_shading(a, fill_color)
    set_cell_border(a, color="CBD5E0", sz="4")

def add_question_row(table, q_idx, q_text, ans_text, widths, is_even=False):
    row = table.add_row()
    cells = row.cells

    # Col 0: TT
    cells[0].width = widths[0]
    p0 = cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(2.5)
    p0.paragraph_format.space_after = Pt(2.5)
    r0 = p0.add_run(str(q_idx))
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(10)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(45, 55, 72)

    # Col 1: Đề bài
    cells[1].width = widths[1]
    p1 = cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p1.paragraph_format.space_before = Pt(2.5)
    p1.paragraph_format.space_after = Pt(2.5)
    p1.paragraph_format.line_spacing = 1.15
    r1 = p1.add_run(q_text)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(10)

    # Col 2: Đáp án đúng
    cells[2].width = widths[2]
    p2 = cells[2].paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_before = Pt(2.5)
    p2.paragraph_format.space_after = Pt(2.5)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(ans_text)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(9.5)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(31, 78, 121)

    # Zebra shading
    if is_even:
        for c in cells:
            set_cell_shading(c, "F7FAFC")

    for c in cells:
        set_cell_border(c, color="CBD5E0", sz="4")

def build_station_content(doc, st):
    # Header Trạm
    p_tram = doc.add_paragraph()
    p_tram.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_tram.paragraph_format.space_before = Pt(6)
    p_tram.paragraph_format.space_after = Pt(2)
    r_tram = p_tram.add_run(f"TRẠM {st['tram_so']}: {st['ten_tram'].upper()} - TUẦN {st['tuan']}")
    r_tram.font.name = "Times New Roman"
    r_tram.font.size = Pt(13.5)
    r_tram.font.bold = True
    r_tram.font.color.rgb = RGBColor(26, 54, 93)

    # Chủ đề & Vật phẩm lưu niệm
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(0)
    p_desc.paragraph_format.space_after = Pt(6)
    
    r_sub1 = p_desc.add_run(f"• Chủ đề trọng tâm: ")
    r_sub1.font.name = "Times New Roman"
    r_sub1.font.size = Pt(10)
    r_sub1.font.bold = True
    r_sub1.font.color.rgb = RGBColor(43, 108, 176)

    r_sub2 = p_desc.add_run(f"{st['chu_de']}\n")
    r_sub2.font.name = "Times New Roman"
    r_sub2.font.size = Pt(10)
    r_sub2.font.italic = True

    if "souvenirs" in st:
        r_souv_lbl = p_desc.add_run(f"• Vật phẩm lưu niệm 3 cấp: ")
        r_souv_lbl.font.name = "Times New Roman"
        r_souv_lbl.font.size = Pt(9.5)
        r_souv_lbl.font.bold = True
        r_souv_lbl.font.color.rgb = RGBColor(180, 83, 9)

        souv = st["souvenirs"]
        r_souv_txt = p_desc.add_run(f"Cấp 1 (50đ): {souv.get('cap1', '')} | Cấp 2 (90đ): {souv.get('cap2', '')} | Cấp 3 (100đ): {souv.get('cap3', '')}\n")
        r_souv_txt.font.name = "Times New Roman"
        r_souv_txt.font.size = Pt(9.5)

    r_sub3 = p_desc.add_run(f"• Ngân hàng đề: ")
    r_sub3.font.name = "Times New Roman"
    r_sub3.font.size = Pt(9.5)
    r_sub3.font.bold = True
    r_sub3.font.color.rgb = RGBColor(197, 48, 48)

    r_sub4 = p_desc.add_run(f"48 câu hỏi chuẩn hóa xoay tua 3 lượt (Vòng 1: 30 câu ghép đôi; Vòng 2: 12 bài điền khuyết; Vòng 3: 6 câu trắc nghiệm Boss)")
    r_sub4.font.name = "Times New Roman"
    r_sub4.font.size = Pt(9.5)

    # Khởi tạo bảng 3 cột: TT | Đề | Đáp án đúng
    widths = [Cm(1.2), Cm(10.8), Cm(4.8)]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["TT", "Đề bài / Câu hỏi", "Đáp án đúng & Hướng dẫn"]
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.width = widths[i]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(h)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(c, "1A365D")
        set_cell_border(c, color="0F233D", sz="6")

    # VÒNG 1: 30 CÂU
    add_section_header_row(
        table,
        "VÒNG 1: KHỞI ĐỘNG (GHÉP ĐÔI KIẾN THỨC - 30 CÂU / 3 LƯỢT XOAY TUA)",
        fill_color="EBF8FF",
        text_color=(43, 108, 176)
    )
    add_sub_group_header_row(table, "▶ LƯỢT THI THỨ NHẤT: Cặp thẻ ghép đôi số 01 đến 10")
    for i in range(10):
        q = st["v1"][i]
        add_question_row(table, i + 1, q["q"], q["a"], widths, is_even=(i % 2 == 1))

    add_sub_group_header_row(table, "▶ LƯỢT THI THỨ HAI: Cặp thẻ ghép đôi số 11 đến 20")
    for i in range(10, 20):
        q = st["v1"][i]
        add_question_row(table, i + 1, q["q"], q["a"], widths, is_even=(i % 2 == 1))

    add_sub_group_header_row(table, "▶ LƯỢT THI THỨ BA: Cặp thẻ ghép đôi số 21 đến 30")
    for i in range(20, 30):
        q = st["v1"][i]
        add_question_row(table, i + 1, q["q"], q["a"], widths, is_even=(i % 2 == 1))

    # VÒNG 2: 12 CÂU
    add_section_header_row(
        table,
        "VÒNG 2: VƯỢT CHƯỚNG NGẠI VẬT (ĐIỀN ĐÁP SỐ TOÁN THỰC TẾ - 12 BÀI / 3 LƯỢT XOAY TUA)",
        fill_color="E6FFFA",
        text_color=(40, 116, 166)
    )
    add_sub_group_header_row(table, "▶ LƯỢT THI THỨ NHẤT: Bài toán thực tế số 31 đến 34")
    for i in range(4):
        q = st["v2"][i]
        add_question_row(table, 31 + i, q["q"], q["a"], widths, is_even=(i % 2 == 1))

    add_sub_group_header_row(table, "▶ LƯỢT THI THỨ HAI: Bài toán thực tế số 35 đến 38")
    for i in range(4, 8):
        q = st["v2"][i]
        add_question_row(table, 31 + i, q["q"], q["a"], widths, is_even=(i % 2 == 1))

    add_sub_group_header_row(table, "▶ LƯỢT THI THỨ BA: Bài toán thực tế số 39 đến 42")
    for i in range(8, 12):
        q = st["v2"][i]
        add_question_row(table, 31 + i, q["q"], q["a"], widths, is_even=(i % 2 == 1))

    # VÒNG 3: 6 CÂU
    add_section_header_row(
        table,
        "VÒNG 3: ĐẤU TRÙM CUỐI (TRẮC NGHIỆM TƯ DUY PHÂN LOẠI MỨC 3 - 6 CÂU / 3 LƯỢT XOAY TUA)",
        fill_color="FFF5F5",
        text_color=(197, 48, 48)
    )
    add_sub_group_header_row(table, "▶ LƯỢT THI THỨ NHẤT: Câu hỏi đấu Boss số 43 đến 44")
    for i in range(2):
        q = st["v3"][i]
        add_question_row(table, 43 + i, q["q"], q["ans_str"], widths, is_even=(i % 2 == 1))

    add_sub_group_header_row(table, "▶ LƯỢT THI THỨ HAI: Câu hỏi đấu Boss số 45 đến 46")
    for i in range(2, 4):
        q = st["v3"][i]
        add_question_row(table, 43 + i, q["q"], q["ans_str"], widths, is_even=(i % 2 == 1))

    add_sub_group_header_row(table, "▶ LƯỢT THI THỨ BA: Câu hỏi đấu Boss số 47 đến 48")
    for i in range(4, 6):
        q = st["v3"][i]
        add_question_row(table, 43 + i, q["q"], q["ans_str"], widths, is_even=(i % 2 == 1))

def generate_hk2_master_document(target_path):
    print(f"Bắt đầu tạo Master File Học kỳ 2: {target_path} ...")
    doc = setup_document_styles("HỌC KỲ 2: TUẦN 19 - 35")

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("BỘ ĐỀ ÔN TẬP TOÁN LỚP 5 - HỌC KỲ 2\nHÀNH TRÌNH THÁM HIỂM XUYÊN VIỆT (TUẦN 19 - 35)")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(17)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(26, 54, 93)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("NGÂN HÀNG 816 CÂU HỎI CHUẨN HÓA (48 CÂU/TRẠM - XOAY TUA 3 LƯỢT)\nKhám phá Tây Nguyên, Duyên hải Nam Trung Bộ, Nam Bộ và Biển Đảo Quê Hương\nDự án: EduBot - BTCT5 (Hành trình thám hiểm Xuyên Việt)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(74, 85, 104)

    summary_table = doc.add_table(rows=1, cols=4)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_widths = [Cm(1.5), Cm(2.2), Cm(6.5), Cm(6.3)]
    s_headers = ["Trạm", "Tuần", "Địa danh thắng cảnh", "Chủ đề trọng tâm"]
    for i, h in enumerate(s_headers):
        c = summary_table.rows[0].cells[i]
        c.width = s_widths[i]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(c, "2B6CB0")

    for st in STATIONS_HK2:
        r = summary_table.add_row().cells
        for i in range(4):
            r[i].width = s_widths[i]
        r[0].paragraphs[0].text = str(st["tram_so"])
        r[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r[1].paragraphs[0].text = f"Tuần {st['tuan']}"
        r[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r[2].paragraphs[0].text = st["ten_tram"]
        r[3].paragraphs[0].text = st["chu_de"]
        for c in r:
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.5)
            set_cell_border(c, color="E2E8F0", sz="4")

    for st in STATIONS_HK2:
        doc.add_page_break()
        build_station_content(doc, st)

    doc.save(target_path)
    print(f"-> Đã lưu thành công Master File HK2 (816 câu): {target_path}")

def generate_grand_master_document(target_path):
    print(f"Bắt đầu tạo ĐẠI TỔNG HỢP 35 TRẠM (1 680 CÂU): {target_path} ...")
    doc = setup_document_styles("TOÀN BỘ 35 TRẠM XUYÊN VIỆT")

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("ĐẠI TỔNG HỢP BỘ ĐỀ TOÁN LỚP 5 - CẢ NĂM HỌC\nHÀNH TRÌNH THÁM HIỂM XUYÊN VIỆT (TRỌN BỘ 35 TRẠM)")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(18)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(26, 54, 93)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("TOÀN BỘ 1 680 CÂU HỎI CHUẨN HÓA (48 CÂU/TRẠM × 35 TRẠM)\nBao phủ 100% Chương trình Giáo dục Phổ thông 2018 môn Toán lớp 5\nDự án: EduBot - BTCT5 (Hành trình thám hiểm Xuyên Việt)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(11.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(74, 85, 104)

    summary_table = doc.add_table(rows=1, cols=4)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_widths = [Cm(1.5), Cm(2.2), Cm(6.5), Cm(6.3)]
    s_headers = ["Trạm", "Tuần", "Địa danh thắng cảnh", "Chủ đề trọng tâm"]
    for i, h in enumerate(s_headers):
        c = summary_table.rows[0].cells[i]
        c.width = s_widths[i]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(c, "1A365D")

    for st in ALL_35_STATIONS:
        r = summary_table.add_row().cells
        for i in range(4):
            r[i].width = s_widths[i]
        r[0].paragraphs[0].text = str(st["tram_so"])
        r[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r[1].paragraphs[0].text = f"Tuần {st['tuan']}"
        r[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r[2].paragraphs[0].text = st["ten_tram"]
        r[3].paragraphs[0].text = st["chu_de"]
        for c in r:
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
            set_cell_border(c, color="E2E8F0", sz="4")

    for st in ALL_35_STATIONS:
        doc.add_page_break()
        build_station_content(doc, st)

    doc.save(target_path)
    print(f"-> Đã lưu thành công GRAND MASTER (1 680 câu): {target_path}")

def generate_hk2_individual_documents(output_dir):
    print(f"Bắt đầu xuất 17 file riêng lẻ HK2 vào: {output_dir} ...")
    for st in STATIONS_HK2:
        doc = setup_document_styles("HỌC KỲ 2: TUẦN 19 - 35")

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(6)
        p_title.paragraph_format.space_after = Pt(2)
        r_t = p_title.add_run(f"ĐỀ KIỂM TRA TOÁN 5 - TUẦN {st['tuan']}")
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(15)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(26, 54, 93)

        build_station_content(doc, st)

        file_name = f"De_Tuan_{st['tuan']:02d}_Tram_{st['tram_so']:02d}.docx"
        file_path = os.path.join(output_dir, file_name)
        doc.save(file_path)
        print(f"   + Đã xuất: {file_name} (48 câu)")

if __name__ == "__main__":
    out_folder = r"c:\Users\Admin\Desktop\EduBot-BTCT5\Đề"
    os.makedirs(out_folder, exist_ok=True)

    hk2_master_path = os.path.join(out_folder, "Bo_De_Toan_5_Tuan_19_den_35.docx")
    generate_hk2_master_document(hk2_master_path)
    generate_hk2_individual_documents(out_folder)

    grand_master_path = os.path.join(out_folder, "Bo_De_Toan_5_Tron_Bo_35_Tuan.docx")
    generate_grand_master_document(grand_master_path)
    print("\nHOÀN TẤT XUẤT TOÀN BỘ 35 TRẠM VÀ HỌC KỲ 2!")
